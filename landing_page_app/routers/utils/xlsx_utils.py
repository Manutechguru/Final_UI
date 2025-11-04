# landing_page_app/routers/utils/xlsx_utils.py
from io import BytesIO
from typing import List, Any
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


# -------------------------------------------
# Safe attribute/dict getter
# -------------------------------------------
def safe_get(row: Any, key: str):
    """Return attribute or dict value safely."""
    try:
        if isinstance(row, dict):
            return row.get(key)
        return getattr(row, key, None)
    except Exception:
        return None


# -------------------------------------------
# Create XLSX (robust header & value resolution)
# -------------------------------------------
def create_xlsx(data: List[Any], columns: List[str], sheet_name: str = "Sheet1") -> BytesIO:
    """
    Creates XLSX file in memory from data list.
    Supports both dicts and SQLAlchemy model objects.
    """

    # --- Canonical header → possible aliases ---
    key_aliases = {
        "Relevant Experience": [
            "relevant_experience", "Relevant Experience", "rel_exp", "relevant_exp"
        ],
        "Skills": [
            "skillset", "Skillset", "Skills", "skills", "skill", "Skill Set",
            "key_skills", "technical_skills"
        ],
    }

    # Ensure columns is a safe copy
    cols = list(columns)

    # If columns contain some skill-like header variants, normalize to exactly "Skills"
    if "Skills" not in cols and any("skill" in str(c).lower() for c in cols):
        cols = ["Skills" if "skill" in str(c).lower() else c for c in cols]

    # Workbook setup
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_font = Font(bold=True)
    center_align = Alignment(horizontal="center", vertical="center")

    # --- Add headers explicitly (guaranteed) ---
    for idx, header in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = header_font
        cell.alignment = center_align

    # --- Value resolver ---
    def resolve_value(row: Any, col_name: str):
        """Find the appropriate value for each column name for a given row."""
        # 1) Try direct match (dict key or attribute)
        v = safe_get(row, col_name)
        if v not in [None, ""]:
            return v

        # 2) Normalized key (spaces -> underscores)
        normalized = col_name.lower().replace(" ", "_")
        v = safe_get(row, normalized)
        if v not in [None, ""]:
            return v

        # 3) Alias list
        if col_name in key_aliases:
            for alt in key_aliases[col_name]:
                v = safe_get(row, alt)
                if v not in [None, ""]:
                    return v

        # 4) fallback blank
        return ""

    # --- Add data rows (starting from row 2) ---
    for row_idx, row in enumerate(data, start=2):
        for col_idx, col in enumerate(cols, start=1):
            val = resolve_value(row, col)
            if isinstance(val, (list, tuple)):
                val = ", ".join(map(str, val))
            # write string or blank (ensure no None)
            ws.cell(row=row_idx, column=col_idx, value="" if val is None else str(val))

    # --- Auto-adjust column widths ---
    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

    # --- Save to in-memory file ---
    file_stream = BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    return file_stream


# -------------------------------------------
# Candidate XLSX Export (keeps compatibility)
# -------------------------------------------
def export_candidates_xlsx(candidates: List[Any]) -> BytesIO:
    """
    Export candidate data to XLSX.
    """
    columns = [
        "Candidate ID", "Candidate Name", "Email", "Contact", "Location",
        "Skills", "Education", "Company", "Relevant Experience",
        "Resume Link", "Status"
    ]
    return create_xlsx(candidates, columns, sheet_name="Candidates")


# -------------------------------------------
# Jobs XLSX Export
# -------------------------------------------
def export_jobs_xlsx(jobs: List[Any]) -> BytesIO:
    columns = ["Job ID", "Job Title", "Client", "Description", "Created At", "Status"]
    return create_xlsx(jobs, columns, sheet_name="Jobs")


# -------------------------------------------
# Clients XLSX Export
# -------------------------------------------
def export_clients_xlsx(clients: List[Any]) -> BytesIO:
    columns = ["Client ID", "Client Name", "Status", "Created At", "Total Jobs"]
    return create_xlsx(clients, columns, sheet_name="Clients")
