# landing_page_app/routers/clients.py
import os
import io
import json
import logging
import traceback
import zipfile
import re
import mimetypes
from datetime import datetime
from typing import Optional, List, Union, Dict, Any
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from urllib.parse import urlparse, unquote

from fastapi import APIRouter, Depends, Request, Form, HTTPException, Body
from fastapi.responses import RedirectResponse, JSONResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session

# XLSX optional support
try:
    from openpyxl import Workbook
except Exception:
    Workbook = None

try:
    import pandas as pd
except Exception:
    pd = None

# Resume parsing optional dependencies
_PYDOCX_AVAILABLE = True
_PYPDF2_AVAILABLE = True
try:
    from docx import Document
except Exception:
    _PYDOCX_AVAILABLE = False
    Document = None

try:
    from PyPDF2 import PdfReader
except Exception:
    _PYPDF2_AVAILABLE = False
    PdfReader = None

# dotenv
from dotenv import load_dotenv
load_dotenv()

from landing_page_app.database import get_db
from landing_page_app.models.clients import Client
from landing_page_app.models.jobs import Job
from landing_page_app.models.candidates import Candidate
from landing_page_app.models.candidate_status_history import CandidateJDMapping

router = APIRouter(prefix="/clients", tags=["Clients"])
templates = Jinja2Templates(directory="landing_page_app/templates")
logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)

# -------------------------
# Gemini / Generative API configuration (via environment)
# -------------------------
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-1.5-flash")
GEMINI_TIMEOUT = int(os.getenv("GEMINI_TIMEOUT", "30"))
GEMINI_MAX_OUTPUT_TOKENS = int(os.getenv("GEMINI_MAX_OUTPUT_TOKENS", "2500"))
GEMINI_CONCURRENCY = int(os.getenv("GEMINI_CONCURRENCY", "3"))
GEMINI_API_BASE = os.getenv("GEMINI_API_BASE", "https://generativelanguage.googleapis.com")

if not GEMINI_API_KEY:
    logger.warning("No Gemini/Google key found in env (GEMINI_API_KEY/GOOGLE_API_KEY). Scoring will fall back to heuristics.")

# Create requests session with retry/backoff
_session = requests.Session()
_retries = Retry(
    total=3,
    backoff_factor=0.6,
    status_forcelist=[429, 500, 502, 503, 504],
    allowed_methods=frozenset(["GET", "POST", "PUT", "DELETE", "HEAD", "OPTIONS"])
)
_adapter = HTTPAdapter(max_retries=_retries)
_session.mount("https://", _adapter)
_session.mount("http://", _adapter)

# -------------------------
# Helpers & exporters
# -------------------------
def _sanitize_filename(s: str) -> str:
    safe = "".join(c for c in (s or "job") if c.isalnum() or c in (" ", "_", "-")).strip()
    return safe.replace(" ", "_") or "job"

XLSX_HEADERS = [
    "Candidate ID", "Name", "Email", "Contact", "Location", "Skills",
    "IT Experience", "Relevant Experience", "Education", "Company",
    "Resume Links", "AI Score", "Comment", "Status"
]

def _normalize_candidate_ids(candidate_ids_input: Optional[Union[str, List[Union[int, str]]]]):
    if candidate_ids_input is None:
        return None
    if isinstance(candidate_ids_input, list):
        ids = []
        for x in candidate_ids_input:
            try:
                ids.append(int(x))
            except Exception:
                continue
        return ids if ids else []
    if isinstance(candidate_ids_input, str):
        s = candidate_ids_input.strip()
        if not s:
            return None
        if s.lower() == "all":
            return None
        parts = [p.strip() for p in s.split(",") if p.strip()]
        ids = []
        for p in parts:
            if p.isdigit():
                ids.append(int(p))
        return ids if ids else []
    return []

def _fetch_candidates_for_job(db: Session, job_id: int, ids_list: Optional[List[int]] = None):
    q = db.query(CandidateJDMapping).filter(CandidateJDMapping.jd_id == job_id)
    if ids_list is not None:
        q = q.filter(CandidateJDMapping.candidate_id.in_(ids_list))
    mappings = q.order_by(CandidateJDMapping.updated_at.desc()).all()
    rows = []
    for mapping in mappings:
        candidate = db.query(Candidate).filter(Candidate.candidates_id == mapping.candidate_id).first()
        if not candidate:
            continue
        rows.append({
            "Candidate ID": candidate.candidates_id,
            "Name": candidate.candidate_name,
            "Email": candidate.email,
            "Contact": candidate.contact,
            "Location": candidate.location,
            "Skills": candidate.skillset,
            "IT Experience": candidate.it_experience,
            "Relevant Experience": candidate.relevant_experience,
            "Education": candidate.education,
            "Company": candidate.company,
            "Resume Links": candidate.resumelinks,
            "AI Score": getattr(mapping, "ai_score", None),
            "Comment": candidate.comment,
            "Status": mapping.stage
        })
    return rows

def _generate_xlsx_bytes(rows: List[dict]):
    if Workbook is not None:
        wb = Workbook()
        ws = wb.active
        ws.append(XLSX_HEADERS)
        for r in rows:
            ws.append([r.get(h, "") if r.get(h, "") is not None else "" for h in XLSX_HEADERS])
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return bio.read()
    elif pd is not None:
        df = pd.DataFrame(rows, columns=XLSX_HEADERS)
        bio = io.BytesIO()
        try:
            df.to_excel(bio, index=False)
        except Exception as e:
            raise HTTPException(status_code=500, detail="Unable to write XLSX. Ensure 'openpyxl' is installed.") from e
        bio.seek(0)
        return bio.read()
    else:
        raise HTTPException(status_code=500, detail="XLSX support requires 'openpyxl' or 'pandas' installed on the server.")

# -------------------------
# JD text fetchers helper
# -------------------------
def _extract_drive_file_id(url: str) -> Optional[str]:
    m = re.search(r"drive\.google\.com/file/d/([^/]+)/?", url)
    if m:
        return m.group(1)
    m = re.search(r"[?&](?:id|fileId)=([^&]+)", url)
    if m:
        return m.group(1)
    return None

def _extract_docs_doc_id(url: str) -> Optional[str]:
    m = re.search(r"docs\.google\.com/document/d/([^/]+)/?", url)
    return m.group(1) if m else None

def try_fetch_jd_text(jd_link: Optional[str]) -> str:
    if not jd_link:
        return ""
    url = jd_link.strip()
    if not url.lower().startswith("http"):
        return url[:4000]

    doc_id = _extract_docs_doc_id(url)
    if doc_id:
        export_url = f"https://docs.google.com/document/d/{doc_id}/export?format=txt"
        try:
            r = _session.get(export_url, timeout=10)
            if r.ok and r.text:
                return r.text.strip()[:4000]
        except Exception as e:
            logger.debug("Docs export failed: %s", e)

    file_id = _extract_drive_file_id(url)
    if file_id:
        try_urls = [
            f"https://drive.google.com/uc?export=download&id={file_id}",
            f"https://drive.usercontent.google.com/download?id={file_id}&export=download"
        ]
        for u in try_urls:
            try:
                r = _session.get(u, timeout=10)
                if r.ok:
                    ctype = (r.headers.get("content-type") or "").lower()
                    if ("text" in ctype) or ("json" in ctype) or ("xml" in ctype) or ("html" in ctype):
                        text = r.text.strip()
                        if text:
                            return text[:4000]
            except Exception as e:
                logger.debug("Drive download attempt failed: %s", e)

    try:
        r = _session.get(url, timeout=8)
        if r.ok and r.text:
            return r.text.strip()[:4000]
    except Exception as e:
        logger.debug("Could not download JD link content: %s", e)

    return url[:4000]

# -------------------------
# Experience parsing & heuristic scorer (fallback)
# -------------------------
def parse_text_experience_to_months(text: Optional[str]) -> Optional[int]:
    if text is None:
        return None
    if isinstance(text, (int, float)):
        try:
            return int(round(float(text) * 12))
        except Exception:
            return None
    s = str(text).strip().lower()
    if s in ("", "n/a", "na", "none", "null"):
        return None
    years = 0.0
    months = 0
    m_year = re.search(r"(\d+(?:\.\d+)?)\s*(?:years|year|yrs|yr)\b", s)
    if m_year:
        try:
            years = float(m_year.group(1))
        except Exception:
            years = 0.0
    m_month = re.search(r"(\d+)\s*(?:months|month|mos|mo)\b", s)
    if m_month:
        try:
            months = int(m_month.group(1))
        except Exception:
            months = 0
    if m_year is None:
        plain_num = re.match(r"^\s*(\d+(?:\.\d+)?)\s*$", s)
        if plain_num:
            try:
                years = float(plain_num.group(1))
            except Exception:
                years = 0.0
    total_months = int(round(years * 12)) + months
    if total_months == 0 and m_year is None and m_month is None and not (plain_num if 'plain_num' in locals() else False):
        return None
    return total_months

def heuristic_score(jd_text: str, cand: Dict[str, Any]) -> Dict[str, Any]:
    try:
        skills_raw = cand.get("skillset") or ""
        skill_tokens = [s.strip().lower() for s in re.split(r"[,;/|]", skills_raw) if s.strip()]
        skill_words = []
        for s in skill_tokens:
            for t in s.split():
                if t:
                    skill_words.append(t)
        skill_set = set(skill_words)

        resume_text = (cand.get("resume_text") or "").lower()
        jd_lower = (jd_text or "").lower()
        combined_search = jd_lower + " " + resume_text

        if not jd_text:
            base = min(80, 30 + len(skill_set) * 6)
            return {
                "score": int(base),
                "explanation": "No JD content available; scoring based on candidate skill count and metadata.",
                "breakdown": {"skill_count": len(skill_set)}
            }

        matched = 0
        for tok in skill_set:
            if tok and tok in combined_search:
                matched += 1
        skill_score = int(round(matched / (len(skill_set) or 1) * 100)) if skill_set else 20

        req_months = None
        m_req = re.search(r"(\d+(?:\.\d+)?)\s*(?:years|yrs|yr)\b", jd_lower)
        if m_req:
            try:
                req_years = float(m_req.group(1))
                req_months = int(round(req_years * 12))
            except Exception:
                req_months = None

        cand_months = None
        for field in ("it_experience", "relevant_experience"):
            cand_txt = cand.get(field)
            if cand_txt:
                cand_months = parse_text_experience_to_months(cand_txt)
                if cand_months is not None:
                    break

        if req_months is not None and cand_months is not None:
            ratio = min(1.0, cand_months / req_months) if req_months > 0 else 1.0
            exp_score = int(round(100 * ratio))
        elif cand_months is not None:
            exp_score = min(90, 30 + int(cand_months / 12) * 10)
        else:
            resume_year_match = re.search(r"(\d+(?:\.\d+)?)\s*(?:years|year|yrs|yr)\b", resume_text)
            if resume_year_match:
                try:
                    res_years = float(resume_year_match.group(1))
                    exp_score = min(90, 30 + int(res_years) * 10)
                except Exception:
                    exp_score = 30
            else:
                exp_score = 30

        final = int(round(0.7 * skill_score + 0.3 * exp_score))
        explanation = f"Skill match {skill_score}%, experience score {exp_score}% (req_months={req_months}, cand_months={cand_months})."
        breakdown = {"skill_score": skill_score, "experience_score": exp_score, "matched_skills": matched, "skill_count": len(skill_set)}
        return {"score": max(0, min(100, final)), "explanation": explanation, "breakdown": breakdown}
    except Exception as e:
        logger.exception("Heuristic scoring failed: %s", e)
        return {"score": 0, "explanation": "Heuristic scoring failed", "breakdown": {}}

# -------------------------
# Gemini wrapper (robust)
# -------------------------
def _extract_text_from_gemini_response(response_json: dict) -> Optional[str]:
    try:
        cands = response_json.get("candidates") or []
        for c in cands:
            content = c.get("content") or {}
            parts = content.get("parts") or []
            texts = []
            for p in parts:
                if isinstance(p, dict) and "text" in p and isinstance(p["text"], str):
                    texts.append(p["text"])
                elif isinstance(p, str):
                    texts.append(p)
            if texts:
                return "".join(texts).strip()
        if isinstance(response_json.get("text"), str):
            return response_json["text"].strip()
    except Exception:
        logger.debug("Error extracting text from Gemini response", exc_info=True)
    return None

def _gemini_call_once(model: str, prompt: str) -> Dict[str, Any]:
    url = f"{GEMINI_API_BASE}/v1beta/models/{model}:generateContent?key={GEMINI_API_KEY}"
    payload = {
        "contents": [
            {
                "role": "user",
                "parts": [{"text": prompt}]
            }
        ],
        "generationConfig": {
            "temperature": 0.0,
            "maxOutputTokens": GEMINI_MAX_OUTPUT_TOKENS
        }
    }
    headers = {"Content-Type": "application/json; charset=utf-8"}
    r = _session.post(url, json=payload, headers=headers, timeout=GEMINI_TIMEOUT)
    r.raise_for_status()
    return r.json()

def call_gemini_for_scoring(jd_text: str, cand: Dict[str, Any]) -> Dict[str, Any]:
    if not GEMINI_API_KEY:
        raise RuntimeError("No Gemini API key available in environment.")

    jd_snippet = (jd_text or "")[:3000]
    prompt_lines = [
        "You are an automated resume-JD matcher. Compare the Job Description, Candidate DB information, and the parsed Resume text.",
        "Produce EXACTLY a JSON object (no surrounding text) with keys: 'score' (integer 0-100),",
        "'explanation' (string), and 'breakdown' (object). The breakdown must include 'skill_match' and 'experience_match' (numbers 0-100).",
        "Return JSON only, nothing else.",
        "",
        "Job Description:",
        jd_snippet,
        "",
        "Candidate Metadata:",
    ]
    for k in ("candidate_id", "candidate_name", "skillset", "location",
              "it_experience", "relevant_experience", "education",
              "company", "resumelinks", "comment"):
        v = cand.get(k) or ""
        prompt_lines.append(f"{k}: {v}")

    resume_text_full = cand.get("resume_text") or ""
    if resume_text_full:
        prompt_lines.append("")
        prompt_lines.append("Resume Text (parsed from CV - truncated):")
        prompt_lines.append(resume_text_full[:3000])

    prompt_lines.append(
        '\nReturn JSON only, for example:\n{"score": 75, "explanation": "...", "breakdown": {"skill_match": 80, "experience_match": 60}}'
    )
    prompt = "\n".join(prompt_lines)

    models_to_try = []
    if GEMINI_MODEL:
        models_to_try.append(GEMINI_MODEL)
    for m in ("gemini-1.5-flash", "gemini-1.5-flash-latest"):
        if m not in models_to_try:
            models_to_try.append(m)

    last_err = None
    for model in models_to_try:
        try:
            j = _gemini_call_once(model, prompt)
            text = _extract_text_from_gemini_response(j)
            if not text:
                raise RuntimeError(f"No text in Gemini response for model={model}: {j}")
            try:
                parsed = json.loads(text)
                if isinstance(parsed, dict) and "score" in parsed:
                    return parsed
                if isinstance(parsed, list):
                    for item in parsed:
                        if isinstance(item, dict) and "score" in item:
                            return item
            except Exception:
                m = re.search(r"(\{[\s\S]*\})", text)
                if m:
                    try:
                        parsed = json.loads(m.group(1))
                        if isinstance(parsed, dict) and "score" in parsed:
                            return parsed
                    except Exception:
                        pass
            raise RuntimeError(f"Gemini returned non-JSON or unexpected content: {repr(text[:800])}")
        except Exception as e:
            last_err = e
            logger.warning("Could not call Gemini with model=%s: %s", model, e)
            continue

    raise RuntimeError(f"Could not call Gemini. Last error: {repr(last_err)}")

# -------------------------
# Resume download & parsing helpers
# -------------------------
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome Safari"

def _guess_ext_from_ctype(ctype: str) -> str:
    if not ctype:
        return ""
    ctype = ctype.split(";")[0].strip().lower()
    mapping = {
        "application/pdf": ".pdf",
        "application/msword": ".doc",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
        "text/plain": ".txt",
        "text/html": ".html",
    }
    if ctype in mapping:
        return mapping[ctype]
    return mimetypes.guess_extension(ctype) or ""

def _parse_cd_filename(cd: str) -> Optional[str]:
    if not cd:
        return None
    m = re.search(r"filename\*=UTF-8''([^;]+)", cd)
    if m:
        try:
            from urllib.parse import unquote as _unquote
            return _unquote(m.group(1))
        except Exception:
            return m.group(1)
    m = re.search(r'filename="([^"]+)"', cd)
    if m:
        return m.group(1)
    m = re.search(r"filename=([^;]+)", cd)
    if m:
        return m.group(1).strip()
    return None

def _drive_direct_urls(file_id: str) -> List[str]:
    return [
        f"https://drive.google.com/uc?export=download&id={file_id}",
        f"https://drive.usercontent.google.com/download?id={file_id}&export=download",
    ]

def _docs_export_urls(doc_id: str) -> List[str]:
    return [
        f"https://docs.google.com/document/d/{doc_id}/export?format=pdf",
        f"https://docs.google.com/document/d/{doc_id}/export?format=docx",
    ]

def _normalize_dropbox(url: str) -> str:
    u = url.replace("www.dropbox.com", "dl.dropboxusercontent.com")
    if "dl=0" in u:
        u = u.replace("dl=0", "dl=1")
    elif "dl=" not in u:
        u = f"{u}&dl=1" if "?" in u else f"{u}?dl=1"
    return u

def _normalize_onedrive(url: str) -> str:
    if "download=" in url:
        return url
    return f"{url}&download=1" if "?" in url else f"{url}?download=1"

def _normalize_resume_link(url: str) -> List[str]:
    if not url:
        return []
    u = url.strip()

    doc_id = _extract_docs_doc_id(u)
    if doc_id:
        return _docs_export_urls(doc_id)

    file_id = _extract_drive_file_id(u)
    if file_id:
        return _drive_direct_urls(file_id)

    if "dropbox.com" in u:
        return [_normalize_dropbox(u)]

    if "1drv.ms" in u or "onedrive.live.com" in u or "sharepoint.com" in u:
        return [_normalize_onedrive(u)]

    return [u]

def _fetch_once(u: str, timeout: int = 25):
    r = _session.get(u, stream=True, timeout=timeout, headers={"User-Agent": UA})
    ctype = (r.headers.get("content-type") or "").lower()

    # Handle Google Drive confirm interstitial pages
    if ("drive.google.com" in u or "drive.usercontent.google.com" in u) and "text/html" in ctype:
        try:
            html = r.text
            m = re.search(r'href="([^"]*?uc\?export=download[^"]*?confirm=[^"&]+[^"]*)"', html)
            if m:
                confirm_url = m.group(1).replace("&amp;", "&")
                if confirm_url.startswith("/"):
                    confirm_url = "https://drive.google.com" + confirm_url
                r = _session.get(confirm_url, stream=True, timeout=timeout, headers={"User-Agent": UA})
                ctype = (r.headers.get("content-type") or "").lower()
        except Exception:
            pass

    r.raise_for_status()

    cd = r.headers.get("content-disposition") or ""
    filename = _parse_cd_filename(cd)

    if not filename:
        try:
            path = urlparse(u).path
            filename = unquote(path.split("/")[-1]) or None
        except Exception:
            filename = None

    base, ext = os.path.splitext(filename or "")
    guessed = _guess_ext_from_ctype(ctype)
    if not ext and guessed:
        filename = (base or "resume") + guessed
    if not filename:
        filename = "resume" + (guessed or "")

    return r.content, filename, ctype

def _fetch_url_bytes(url: str, timeout: int = 25):
    attempts = _normalize_resume_link(url)
    last_exc = None
    for u in attempts:
        try:
            content, fname, ctype = _fetch_once(u, timeout=timeout)
            return content, fname, ctype
        except Exception as e:
            last_exc = e
            continue
    if last_exc:
        raise last_exc
    raise RuntimeError("Unable to fetch resume: no valid URLs to try")

def _is_probably_valid_file(content: bytes, filename: str, ctype: str) -> bool:
    if not content:
        return False
    if isinstance(content, str):
        content_bytes = content.encode("utf-8", errors="ignore")
    else:
        content_bytes = content

    low = content_bytes[:512].lower()
    if content_bytes.startswith(b"%PDF"):
        return True
    if content_bytes.startswith(b"PK"):
        return True
    if b"\x00" in content_bytes[:1024]:
        return True
    if ctype and not ctype.startswith("text/") and "html" not in (ctype or ""):
        return True
    if filename and filename.lower().endswith(".txt"):
        return True
    if b"<html" in low or b"<!doctype html" in low or b"<head" in low:
        return False
    try:
        s = content_bytes[:400].decode("utf-8", errors="ignore").lower()
        if "error" in s or "unauthorized" in s or "forbidden" in s or "401" in s:
            return False
    except Exception:
        pass
    return True

def _extract_text_from_file_bytes(content: bytes, filename: str) -> str:
    text = ""
    try:
        fn = (filename or "").lower()
        if fn.endswith(".pdf"):
            if not _PYPDF2_AVAILABLE:
                logger.warning("PyPDF2 not installed; cannot parse PDF resume '%s'.", filename)
                return ""
            try:
                reader = PdfReader(io.BytesIO(content))
                parts = []
                for page in reader.pages:
                    try:
                        page_text = page.extract_text() or ""
                        parts.append(page_text)
                    except Exception:
                        continue
                text = "\n".join(parts)
            except Exception as e:
                logger.debug("Pdf parsing failed for %s: %s", filename, e)
                text = ""
        elif fn.endswith(".docx") or fn.endswith(".doc"):
            if not _PYDOCX_AVAILABLE:
                logger.warning("python-docx not installed; cannot parse DOCX resume '%s'.", filename)
                return ""
            try:
                doc = Document(io.BytesIO(content))
                text = "\n".join([p.text for p in doc.paragraphs])
            except Exception as e:
                logger.debug("Docx parsing failed for %s: %s", filename, e)
                text = ""
        else:
            try:
                s = content.decode("utf-8", errors="ignore")
                if len(s.strip()) > 10:
                    text = s
            except Exception:
                text = ""
    except Exception as outer:
        logger.exception("Unexpected resume extraction error for %s: %s", filename, outer)
        text = ""
    return (text or "").strip()[:4000]

# -------------------------
# Download resumes (ZIP)
# -------------------------
@router.post("/{client_id}/job/{job_id}/download-resumes")
def download_resumes_json(client_id: int, job_id: int, payload: dict = Body(...), db: Session = Depends(get_db)):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    ids_norm = _normalize_candidate_ids(payload.get("candidate_ids"))
    if ids_norm == []:
        raise HTTPException(status_code=400, detail="No valid candidate ids provided")

    if ids_norm is None:
        mappings = db.query(CandidateJDMapping).filter(CandidateJDMapping.jd_id == job_id).all()
        ids_to_fetch = [m.candidate_id for m in mappings]
    else:
        ids_to_fetch = ids_norm

    if not ids_to_fetch:
        raise HTTPException(status_code=400, detail="No candidates available for the requested job.")

    buf = io.BytesIO()
    zf = zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED)
    missing = []

    for cid in ids_to_fetch:
        cand = db.query(Candidate).filter(Candidate.candidates_id == cid).first()
        if not cand:
            missing.append(f"{cid}: Candidate not found")
            continue

        links_raw = (cand.resumelinks or "").strip()
        links = [l.strip() for l in re.split(r"[\n,;]+", links_raw) if l.strip()]
        if not links:
            missing.append(f"{cid} ({cand.candidate_name}): no resume link")
            continue

        got_file = False
        for idx, link in enumerate(links, start=1):
            try:
                content, fname, ctype = _fetch_url_bytes(link)
                if isinstance(content, str):
                    content = content.encode("utf-8", errors="ignore")

                if not _is_probably_valid_file(content, fname, ctype):
                    raise RuntimeError(f"Fetched content for {link} looks like HTML/error (content-type={ctype})")

                # sanitize filename and ensure extension
                safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", fname or "")
                if not safe_name:
                    base = f"{(cand.candidate_name or 'candidate').replace(' ', '_')}_{cid}_{idx}"
                    safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", base) + (_guess_ext_from_ctype(ctype) or ".pdf")

                base2, ext2 = os.path.splitext(safe_name)
                if not ext2:
                    ext2 = _guess_ext_from_ctype(ctype) or ".pdf"
                    safe_name = base2 + ext2

                arc = safe_name
                i = 1
                existing = {info.filename for info in zf.infolist()}
                while arc in existing:
                    arc = f"{base2}_{i}{ext2}"
                    i += 1

                zf.writestr(arc, content)
                got_file = True
                break
            except Exception as e:
                logger.warning("Failed to fetch resume for candidate %s url=%s: %s", cid, link, e)
                missing.append(
                    f"{cid} ({cand.candidate_name}): Could not download from {link} -> {e}. "
                    f"Tip: For Google Drive, set sharing to 'Anyone with the link can view'."
                )
                continue

        if not got_file:
            continue

    if missing:
        zf.writestr("missing_resumes.txt", "\n".join(missing))

    zf.close()
    buf.seek(0)
    filename = f"{_sanitize_filename(job.job_title)}_resumes_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.zip"

    return StreamingResponse(
        buf,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

# -------------------------
# Worker used in parallel scoring (does NOT touch DB)
# -------------------------
def _score_candidate_worker(cand: Dict[str, Any], jd_text: str, explain_flag: bool) -> Dict[str, Any]:
    cid = cand.get("candidate_id")
    try:
        cand = dict(cand)  # shallow copy
        cand.setdefault("resume_text", "")

        links_raw = (cand.get("resumelinks") or "").strip()
        links = [l.strip() for l in re.split(r"[\n,;]+", links_raw) if l.strip()]
        if links:
            for link in links:
                try:
                    content, fname, ctype = _fetch_url_bytes(link, timeout=20)
                    if isinstance(content, str):
                        content = content.encode("utf-8", errors="ignore")
                    if not _is_probably_valid_file(content, fname, ctype):
                        raise RuntimeError("Fetched content looks invalid (HTML/error), skipping.")
                    parsed = _extract_text_from_file_bytes(content, fname)
                    if parsed:
                        cand["resume_text"] = parsed
                        break
                except Exception as e:
                    logger.debug("Worker: could not fetch/parse resume for candidate %s (%s): %s", cid, link, e)
                    continue

        ai_out = None
        if GEMINI_API_KEY:
            try:
                ai_out = call_gemini_for_scoring(jd_text, cand)
            except Exception as e:
                logger.warning("Worker: AI scoring failed for candidate %s: %s", cid, e)
                logger.debug(traceback.format_exc())
                ai_out = None

        if ai_out and isinstance(ai_out, dict) and "score" in ai_out:
            score_val = int(ai_out.get("score", 0))
            explanation = ai_out.get("explanation", "") if explain_flag else ""
            breakdown = ai_out.get("breakdown", {}) if explain_flag else {}
            return {"candidate_id": cid, "score": max(0, min(100, score_val)), "explanation": explanation, "breakdown": breakdown}
        else:
            fallback = heuristic_score(jd_text, cand)
            return {
                "candidate_id": cid,
                "score": int(fallback.get("score", 0)),
                "explanation": fallback.get("explanation", "") if explain_flag else "",
                "breakdown": fallback.get("breakdown", {}) if explain_flag else {}
            }
    except Exception as e:
        logger.exception("Worker scoring error for candidate %s: %s", cid, e)
        return {"error": {"candidate_id": cid, "error": str(e)}}

# -------------------------
# Scoring endpoint (AI + fallback) with DB persistence of ai_score
# -------------------------
@router.post("/{client_id}/job/{job_id}/score-candidates")
def score_candidates_endpoint(
    client_id: int,
    job_id: int,
    payload: dict = Body(...),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    raw_ids = payload.get("candidate_ids", "all")
    ids = _normalize_candidate_ids(raw_ids)
    if ids == []:
        raise HTTPException(status_code=400, detail="No valid candidate ids provided")

    explain_flag = bool(payload.get("explain", True))

    if getattr(job, "job_description", None):
        jd_text = try_fetch_jd_text(job.job_description)
    else:
        jd_text = job.job_title or ""

    scores = []
    errors = []

    # Resolve candidate ids to actual candidates (snapshot) BEFORE threading
    if ids is None:
        mappings = db.query(CandidateJDMapping).filter(CandidateJDMapping.jd_id == job_id).all()
        ids_iter = [m.candidate_id for m in mappings]
    else:
        ids_iter = ids

    candidate_snapshots = []
    for cid in ids_iter:
        cand_obj = db.query(Candidate).filter(Candidate.candidates_id == cid).first()
        if not cand_obj:
            errors.append({"candidate_id": cid, "error": "Candidate not found"})
            continue
        cand_snapshot = {
            "candidate_id": cand_obj.candidates_id,
            "candidate_name": cand_obj.candidate_name,
            "skillset": cand_obj.skillset,
            "location": cand_obj.location,
            "it_experience": cand_obj.it_experience,
            "relevant_experience": cand_obj.relevant_experience,
            "education": cand_obj.education,
            "company": cand_obj.company,
            "resumelinks": cand_obj.resumelinks,
            "comment": cand_obj.comment,
        }
        candidate_snapshots.append((cid, cand_snapshot))

    max_workers = max(1, min(len(candidate_snapshots), GEMINI_CONCURRENCY))
    with ThreadPoolExecutor(max_workers=max_workers) as ex:
        futures = {ex.submit(_score_candidate_worker, snap, jd_text, explain_flag): cid for cid, snap in candidate_snapshots}
        for fut in as_completed(futures):
            cid_for_future = futures[fut]
            try:
                res = fut.result()
            except Exception as e:
                logger.exception("Thread error for candidate %s: %s", cid_for_future, e)
                errors.append({"candidate_id": cid_for_future, "error": str(e)})
                continue

            if not res:
                errors.append({"candidate_id": cid_for_future, "error": "No response"})
                continue

            if "error" in res:
                errors.append(res["error"])
                continue

            # Persist AI score to candidate_jd_mapping.ai_score
            try:
                score_value = int(res.get("score", 0))
            except Exception:
                score_value = 0

            try:
                mapping = db.query(CandidateJDMapping).filter(
                    CandidateJDMapping.candidate_id == cid_for_future,
                    CandidateJDMapping.jd_id == job_id
                ).first()
                if mapping:
                    mapping.ai_score = score_value
                    try:
                        mapping.updated_at = datetime.utcnow()
                    except Exception:
                        pass
                    db.add(mapping)
                    db.commit()
                else:
                    new_map = CandidateJDMapping(
                        candidate_id=cid_for_future,
                        jd_id=job_id,
                        stage="Applied",
                        updated_at=datetime.utcnow(),
                        ai_score=score_value
                    )
                    db.add(new_map)
                    db.commit()
            except Exception as e:
                logger.exception("Failed to persist ai_score for candidate %s job %s: %s", cid_for_future, job_id, e)
                try:
                    db.rollback()
                except Exception:
                    pass
                errors.append({"candidate_id": cid_for_future, "error": f"Failed to persist ai_score: {e}"})
                scores.append(res)
                continue

            scores.append(res)

    return JSONResponse(content={"scores": scores, "errors": errors})

# -------------------------
# Reset AI score endpoints
# -------------------------
@router.post("/{client_id}/job/{job_id}/reset-ai-score")
def reset_ai_scores_for_job_or_candidates(
    client_id: int,
    job_id: int,
    payload: dict = Body(...),
    db: Session = Depends(get_db)
):
    """
    Reset AI score(s). Payload:
      { "candidate_ids": "1,2,3" }  -> reset listed
      { "candidate_ids": ["1","2"] } -> reset listed
      { "candidate_ids": "all" } -> reset all candidates for the job
    """
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    raw_ids = payload.get("candidate_ids", "all")
    ids = _normalize_candidate_ids(raw_ids)

    if ids == []:
        raise HTTPException(status_code=400, detail="No valid candidate ids provided")

    if ids is None:
        # reset all mappings for this job
        mappings = db.query(CandidateJDMapping).filter(CandidateJDMapping.jd_id == job_id).all()
        count = 0
        for m in mappings:
            if getattr(m, "ai_score", None) is not None:
                m.ai_score = None
                try:
                    m.updated_at = datetime.utcnow()
                except Exception:
                    pass
                db.add(m)
                count += 1
        db.commit()
        return JSONResponse(content={"reset_count": count})
    else:
        reset_count = 0
        for cid in ids:
            mapping = db.query(CandidateJDMapping).filter(CandidateJDMapping.jd_id == job_id, CandidateJDMapping.candidate_id == cid).first()
            if mapping and getattr(mapping, "ai_score", None) is not None:
                mapping.ai_score = None
                try:
                    mapping.updated_at = datetime.utcnow()
                except Exception:
                    pass
                db.add(mapping)
                reset_count += 1
        db.commit()
        return JSONResponse(content={"reset_count": reset_count})

@router.post("/{client_id}/job/{job_id}/reset-ai-score/{candidate_id}")
def reset_single_candidate_ai_score(client_id: int, job_id: int, candidate_id: int, db: Session = Depends(get_db)):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    mapping = db.query(CandidateJDMapping).filter(CandidateJDMapping.jd_id == job_id, CandidateJDMapping.candidate_id == candidate_id).first()
    if not mapping:
        raise HTTPException(status_code=404, detail="Mapping not found")
    mapping.ai_score = None
    try:
        mapping.updated_at = datetime.utcnow()
    except Exception:
        pass
    db.add(mapping)
    db.commit()
    return JSONResponse(content={"candidate_id": candidate_id, "reset": True})

# -------------------------
# Pages & CRUD (returns ai_score for templates)
# -------------------------
@router.get("/client_jobs")
def redirect_client_jobs():
    return RedirectResponse(url="/clients/new-arrivals")

@router.get("/new-arrivals")
def new_arrivals_page(request: Request, db: Session = Depends(get_db), message: str = ""):
    clients = db.query(Client).order_by(Client.created_at.desc()).all()
    active_count = db.query(Client).filter(Client.status == "active").count()
    inactive_count = db.query(Client).filter(Client.status == "inactive").count()
    return templates.TemplateResponse(
        "new_arrivals.html",
        {
            "request": request,
            "clients": clients,
            "active_count": active_count,
            "inactive_count": inactive_count,
            "message": message,
        },
    )

@router.post("/new-arrivals")
def add_new_client(request: Request, client_name: str = Form(...), db: Session = Depends(get_db)):
    client_name = client_name.strip()
    message = ""
    if not client_name:
        message = "Client name is required!"
    else:
        existing_client = db.query(Client).filter(Client.client_name == client_name).first()
        if existing_client:
            message = f"Client '{client_name}' already exists!"
        else:
            new_client = Client(client_name=client_name, created_at=datetime.utcnow(), status="active")
            db.add(new_client)
            db.commit()
            db.refresh(new_client)
            message = f"Client '{client_name}' added successfully!"

    clients = db.query(Client).order_by(Client.created_at.desc()).all()
    active_count = db.query(Client).filter(Client.status == "active").count()
    inactive_count = db.query(Client).filter(Client.status == "inactive").count()

    return templates.TemplateResponse(
        "new_arrivals.html",
        {
            "request": request,
            "clients": clients,
            "active_count": active_count,
            "inactive_count": inactive_count,
            "message": message,
        },
    )

@router.post("/delete/{client_id}")
def delete_client(client_id: int, db: Session = Depends(get_db)):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")

    db.query(Job).filter(Job.client_id == client_id).delete(synchronize_session=False)
    db.delete(client)
    db.commit()
    return JSONResponse(content={"message": f"Client '{client.client_name}' and all its jobs deleted successfully!"})

@router.post("/toggle-client/{client_id}")
@router.post("/toggle-active/{client_id}")
def toggle_client_status(client_id: int, db: Session = Depends(get_db)):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")

    client.status = "inactive" if client.status == "active" else "active"
    db.commit()
    db.refresh(client)
    return {"client_id": client.client_id, "new_status": client.status}

@router.get("/stats")
def get_client_job_stats(db: Session = Depends(get_db)):
    total_clients = db.query(Client).count()
    total_jobs = db.query(Job).count()
    active_jobs = db.query(Job).filter(Job.status == "active").count()
    inactive_jobs = db.query(Job).filter(Job.status == "inactive").count()
    active_clients = db.query(Client).filter(Client.status == "active").count()
    inactive_clients = db.query(Client).filter(Client.status == "inactive").count()

    return JSONResponse(
        content={
            "total_clients": total_clients,
            "active_clients": active_clients,
            "inactive_clients": inactive_clients,
            "total_jobs": total_jobs,
            "active_jobs": active_jobs,
            "inactive_jobs": inactive_jobs,
        }
    )

@router.get("/{client_id}")
def view_client_jobs(request: Request, client_id: int, db: Session = Depends(get_db), message: str = ""):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        return templates.TemplateResponse("error.html", {"request": request, "message": "Client not found!"})

    jobs = db.query(Job).filter(Job.client_id == client_id).order_by(Job.created_at.desc()).all()
    return templates.TemplateResponse(
        "client_jobs.html", {"request": request, "client": client, "jobs": jobs, "message": message}
    )

@router.post("/{client_id}/add-job")
def add_job_for_client(request: Request, client_id: int, job_title: str = Form(...), jd_link: str = Form(...), db: Session = Depends(get_db)):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        return templates.TemplateResponse("error.html", {"request": request, "message": "Client not found!"})

    job_title = job_title.strip()
    jd_link = jd_link.strip()
    message = ""

    if not job_title:
        message = "Job title is required!"
    elif not jd_link:
        message = "JD link is required!"
    else:
        new_job = Job(
            client_id=client_id,
            job_title=job_title,
            job_description=jd_link,
            status="active",
            created_at=datetime.utcnow(),
        )
        db.add(new_job)
        db.commit()
        db.refresh(new_job)
        message = f"Job '{job_title}' added successfully under client '{client.client_name}'!"

    jobs = db.query(Job).filter(Job.client_id == client_id).order_by(Job.created_at.desc()).all()
    return templates.TemplateResponse(
        "client_jobs.html", {"request": request, "client": client, "jobs": jobs, "message": message}
    )

@router.post("/{client_id}/delete-job/{job_id}")
def delete_job(client_id: int, job_id: int, db: Session = Depends(get_db)):
    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    db.query(CandidateJDMapping).filter(CandidateJDMapping.jd_id == job_id).delete(synchronize_session=False)
    db.delete(job)
    db.commit()
    return JSONResponse(content={"message": f"Job '{job.job_title}' deleted successfully!"})

@router.post("/{client_id}/toggle-job/{job_id}")
def toggle_job_status(request: Request, client_id: int, job_id: int, db: Session = Depends(get_db)):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")

    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    job.status = "inactive" if job.status == "active" else "active"
    db.commit()
    db.refresh(job)

    jobs = db.query(Job).filter(Job.client_id == client_id).order_by(Job.created_at.desc()).all()
    return templates.TemplateResponse(
        "client_jobs.html", {"request": request, "client": client, "jobs": jobs, "message": ""}
    )

@router.get("/{client_id}/job/{job_id}/candidates")
def view_candidates_for_job(request: Request, client_id: int, job_id: int, db: Session = Depends(get_db)):
    """
    NOTE: This function includes mapping.ai_score so templates will always receive persisted scores.
    Ensure your template reads 'ai_score' from each candidate dict (it is provided below).
    """
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        return templates.TemplateResponse("error.html", {"request": request, "message": "Client not found!"})

    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        return templates.TemplateResponse("error.html", {"request": request, "message": "Job not found!"})

    candidate_mappings = db.query(CandidateJDMapping).filter(CandidateJDMapping.jd_id == job_id).order_by(
        CandidateJDMapping.updated_at.desc()
    ).all()

    candidates = []
    for mapping in candidate_mappings:
        candidate = db.query(Candidate).filter(Candidate.candidates_id == mapping.candidate_id).first()
        if candidate:
            candidates.append(
                {
                    "candidates_id": candidate.candidates_id,
                    "candidate_name": candidate.candidate_name,
                    "email": candidate.email,
                    "contact": candidate.contact,
                    "location": candidate.location,
                    "skillset": candidate.skillset,
                    "relevant_experience": candidate.relevant_experience,
                    "it_experience": candidate.it_experience,
                    "education": candidate.education,
                    "company": candidate.company,
                    "resumelinks": candidate.resumelinks,
                    "comment": candidate.comment,
                    "status": mapping.stage,
                    "ai_score": getattr(mapping, "ai_score", None)
                }
            )

    return templates.TemplateResponse(
        "job_candidates.html", {"request": request, "client": client, "job": job, "candidates": candidates}
    )

@router.post("/{client_id}/job/{job_id}/remove-candidates")
def remove_candidates_from_job(request: Request, client_id: int, job_id: int, candidate_ids: str = Form(...), db: Session = Depends(get_db)):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")

    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    ids_list = [int(cid.strip()) for cid in candidate_ids.split(",") if cid.strip().isdigit()]
    db.query(CandidateJDMapping).filter(
        CandidateJDMapping.jd_id == job_id, CandidateJDMapping.candidate_id.in_(ids_list)
    ).delete(synchronize_session=False)

    db.commit()
    return JSONResponse(content={"message": f"Removed {len(ids_list)} candidate(s) from job."})

# Downloads (XLSX)
@router.post("/{client_id}/job/{job_id}/download-candidates")
def download_candidates_post(
    client_id: int,
    job_id: int,
    candidate_ids: str = Form(...),
    file_format: str = Form("xlsx"),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    ids_norm = _normalize_candidate_ids(candidate_ids)
    if ids_norm == []:
        raise HTTPException(status_code=400, detail="No valid candidate ids provided")

    rows = _fetch_candidates_for_job(db, job_id, ids_norm)
    fmt = (file_format or "xlsx").lower()

    if fmt != "xlsx":
        raise HTTPException(status_code=400, detail="Only 'xlsx' format is supported now.")

    filename = f"{_sanitize_filename(job.job_title)}_candidates_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.xlsx"
    data_bytes = _generate_xlsx_bytes(rows)
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    return StreamingResponse(io.BytesIO(data_bytes),
                             media_type=media_type,
                             headers={"Content-Disposition": f'attachment; filename=\"{filename}\"'})

@router.post("/{client_id}/job/{job_id}/download")
def download_candidates_json(
    client_id: int,
    job_id: int,
    payload: dict = Body(...),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    ids_norm = _normalize_candidate_ids(payload.get("candidate_ids"))
    if ids_norm == []:
        raise HTTPException(status_code=400, detail="No valid candidate ids provided")

    rows = _fetch_candidates_for_job(db, job_id, ids_norm)
    fmt = (payload.get("file_format") or "xlsx").lower()
    if fmt != "xlsx":
        raise HTTPException(status_code=400, detail="Only 'xlsx' format is supported now.")

    filename = f"{_sanitize_filename(job.job_title)}_candidates_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.xlsx"
    data_bytes = _generate_xlsx_bytes(rows)
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    return StreamingResponse(io.BytesIO(data_bytes),
                             media_type=media_type,
                             headers={"Content-Disposition": f'attachment; filename=\"{filename}\"'})

# Candidate details + bulk update
@router.get("/{client_id}/job/{job_id}/candidate/{candidate_id}")
def get_candidate_details(client_id: int, job_id: int, candidate_id: int, db: Session = Depends(get_db)):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    mapping = db.query(CandidateJDMapping).filter(CandidateJDMapping.jd_id == job_id, CandidateJDMapping.candidate_id == candidate_id).first()
    if not mapping:
        raise HTTPException(status_code=404, detail="Candidate not linked to this job")
    cand = db.query(Candidate).filter(Candidate.candidates_id == candidate_id).first()
    if not cand:
        raise HTTPException(status_code=404, detail="Candidate not found")
    data = {
        "candidates_id": cand.candidates_id,
        "candidate_name": cand.candidate_name,
        "email": cand.email,
        "contact": cand.contact,
        "location": cand.location,
        "skillset": cand.skillset,
        "it_experience": cand.it_experience,
        "relevant_experience": cand.relevant_experience,
        "education": cand.education,
        "company": cand.company,
        "resumelinks": cand.resumelinks,
        "comment": cand.comment,
        "notice_period": getattr(cand, "notice_period", None),
        "status": mapping.stage,
        "ai_score": getattr(mapping, "ai_score", None)
    }
    return JSONResponse(content={"candidate": data})

def _apply_candidate_updates(db: Session, cand: Candidate, updates: dict) -> dict:
    allowed = ["candidate_name", "email", "contact", "location", "skillset", "it_experience",
               "relevant_experience", "education", "company", "resumelinks", "comment", "notice_period"]
    changed = False
    for k in allowed:
        if k in updates:
            new_val = updates.get(k)
            if isinstance(new_val, str) and new_val.strip() == "":
                new_val = None
            if getattr(cand, k, None) != new_val:
                setattr(cand, k, new_val)
                changed = True
    if changed:
        try:
            cand.updated_at = datetime.utcnow()
        except Exception:
            pass
    return {
        "candidates_id": cand.candidates_id,
        "candidate_name": cand.candidate_name,
        "email": cand.email,
        "contact": cand.contact,
        "location": cand.location,
        "skillset": cand.skillset,
        "it_experience": cand.it_experience,
        "relevant_experience": cand.relevant_experience,
        "education": cand.education,
        "company": cand.company,
        "resumelinks": cand.resumelinks,
        "comment": cand.comment,
        "notice_period": getattr(cand, "notice_period", None)
    }

@router.post("/{client_id}/job/{job_id}/update-candidates")
def update_candidates_bulk(client_id: int, job_id: int, payload: dict = Body(...), db: Session = Depends(get_db)):
    client = db.query(Client).filter(Client.client_id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    job = db.query(Job).filter(Job.job_id == job_id, Job.client_id == client_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    candidates_payload = payload.get("candidates")
    if not isinstance(candidates_payload, list) or not candidates_payload:
        raise HTTPException(status_code=400, detail="Payload must include 'candidates' array with at least one item")

    updated = []
    errors = []

    for item in candidates_payload:
        try:
            cid = item.get("candidates_id") or item.get("candidate_id") or item.get("id")
            if not cid:
                errors.append({"candidate_id": None, "error": "Missing candidate id in payload item"})
                continue
            try:
                cid = int(cid)
            except Exception:
                errors.append({"candidate_id": cid, "error": "Invalid candidate id"})
                continue
            mapping = db.query(CandidateJDMapping).filter(CandidateJDMapping.jd_id == job_id, CandidateJDMapping.candidate_id == cid).first()
            if not mapping:
                errors.append({"candidate_id": cid, "error": "Candidate not linked to this job"})
                continue
            cand = db.query(Candidate).filter(Candidate.candidates_id == cid).first()
            if not cand:
                errors.append({"candidate_id": cid, "error": "Candidate not found"} )
                continue
            result_dict = _apply_candidate_updates(db, cand, item)
            db.add(cand)
            updated.append(result_dict)
        except Exception as e:
            logger.exception("Error updating candidate %s: %s", item.get("candidates_id"), e)
            errors.append({"candidate_id": item.get("candidates_id"), "error": str(e)})

    try:
        db.commit()
    except Exception as e:
        logger.exception("Commit failed after candidate updates: %s", e)
        db.rollback()
        raise HTTPException(status_code=500, detail="Failed to save candidate updates")

    return JSONResponse(content={"updated": updated, "errors": errors})

# End of file
